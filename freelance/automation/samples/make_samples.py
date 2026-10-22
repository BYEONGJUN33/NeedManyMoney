# 자동화 도구 시연용 가짜 데이터 만들기: python3 make_samples.py
# - 매출_강남점/홍대점/판교점.xlsx: 위에 제목 줄이 있고, 지점마다 열 순서·열 이름이 조금씩 다름
# - 스마트스토어_주문내역_샘플.xlsx: 같은 사람이 여러 개 주문한 경우 포함 (합포장 확인용)
import datetime as dt
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
rnd = random.Random(42)
ITEMS = [('아메리카노', 4500), ('카페라떼', 5000), ('바닐라라떼', 5500), ('콜드브루', 5000), ('치즈케이크', 6500), ('원두 200g', 16000)]


def sales(branch, order, extra=None, rename=None):
    wb = Workbook()
    ws = wb.active
    ws.title = '3분기'
    ws.append([f'2026년 3분기 매출 — {branch}'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    header = list(order) + ([extra] if extra else [])
    ws.append([rename.get(h, h) if rename else h for h in header])
    total = 0
    day = dt.date(2026, 7, 1)
    while day <= dt.date(2026, 9, 30):
        for _ in range(rnd.randint(0, 2)):
            name, price = rnd.choice(ITEMS)
            qty = rnd.randint(3, 40)
            row = {'일자': day, '품목': name, '수량': qty, '단가': price, '매출액': qty * price, '담당자': rnd.choice(['민지', '도윤', '서연'])}
            total += qty * price
            ws.append([row[h] for h in header])
        day += dt.timedelta(days=rnd.randint(1, 3))
    for cell in ws['A' if order[0] == '일자' else 'B'][3:]:
        cell.number_format = 'yyyy-mm-dd'
    ws.append([])
    ws.append(['합계', '', '', '', total])
    wb.save(os.path.join(HERE, f'매출_{branch}.xlsx'))


sales('강남점', ['일자', '품목', '수량', '단가', '매출액'])
sales('홍대점', ['일자', '품목', '단가', '수량', '매출액'], extra='담당자')
sales('판교점', ['품목', '일자', '수량', '단가', '매출액'], rename={'매출액': '매출 금액'})

# 주문 내역 (스마트스토어 '발주/발송 확인' 내려받기와 비슷한 열 구성)
PRODUCTS = [('콜롬비아 수프리모 원두', ['200g / 홀빈', '200g / 핸드드립용 분쇄', '500g / 홀빈']), ('드립백 선물세트', ['10개입', '20개입']), ('스테인리스 드리퍼', ['1~2인용', '2~4인용']), ('디카페인 원두', ['200g / 홀빈'])]
PEOPLE = [('김하늘', '010-1111-2222', '06236', '서울특별시 강남구 테헤란로 000 101동 1001호'), ('이도윤', '010-3333-4444', '04039', '서울특별시 마포구 양화로 000 3층'),
          ('박서연', '010-5555-6666', '13494', '경기도 성남시 분당구 판교역로 000 A동 702호'), ('최지우', '010-7777-8888', '48058', '부산광역시 해운대구 센텀중앙로 000 1203호'),
          ('정민준', '010-9999-0000', '35209', '대전광역시 서구 둔산로 000 5층'), ('강예린', '010-1212-3434', '61947', '광주광역시 서구 상무중앙로 000 802호'),
          ('윤시우', '010-5656-7878', '21999', '인천광역시 연수구 센트럴로 000 1105호'), ('한지민', '010-9090-1212', '63309', '제주특별자치도 제주시 첨단로 000')]
wb = Workbook()
ws = wb.active
ws.title = '발주발송관리'
ws.append(['상품주문번호', '주문번호', '주문일시', '구매자명', '수취인명', '수취인연락처1', '우편번호', '통합배송지', '상품명', '옵션정보', '수량', '배송메세지'])
no = 2026102200000
t = dt.datetime(2026, 10, 21, 9, 0)
for i in range(26):
    who = PEOPLE[i % len(PEOPLE)] if i < 18 else rnd.choice(PEOPLE[:3])
    order_no = 202610210000 + (i if i < 18 else i - 18)
    prod, opts = rnd.choice(PRODUCTS)
    no += 1
    t += dt.timedelta(minutes=rnd.randint(5, 50))
    ws.append([str(no), str(order_no), t.strftime('%Y-%m-%d %H:%M'), who[0], who[0], who[1], who[2], who[3], prod, rnd.choice(opts), rnd.randint(1, 3), rnd.choice(['', '', '문 앞에 놓아 주세요', '경비실에 맡겨 주세요', '배송 전 연락 부탁드립니다'])])
wb.save(os.path.join(HERE, '스마트스토어_주문내역_샘플.xlsx'))
print('ok')
